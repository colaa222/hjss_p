from openpyxl import load_workbook


def read_targets_from_excel(file_path: str, sheet_name: str | None = None) -> list[dict]:
    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name else wb.active

    rows = []
    row_no = 2  # 1행은 헤더

    while True:
        target = ws[f"A{row_no}"].value
        message = ws[f"B{row_no}"].value

        if target is None and message is None:
            break

        rows.append(
            {
                "excel_row": row_no,
                "target": "" if target is None else str(target).strip(),
                "message": "" if message is None else str(message).strip(),
            }
        )
        row_no += 1

    return rows


def write_results_to_excel(
    file_path: str,
    results: list[dict],
    status_col: str = "C",
    detail_col: str = "D",
    sheet_name: str | None = None,
) -> None:
    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name else wb.active

    for item in results:
        row_no = item.get("excel_row")
        if not row_no:
            continue

        ws[f"{status_col}{row_no}"] = 1 if item.get("ok") else 0
        ws[f"{detail_col}{row_no}"] = item.get("message_detail", "")

    wb.save(file_path)